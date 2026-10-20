"""Integration checks. Requires PowerShell 7, Poppler, Tesseract, Pillow, reportlab."""
import gzip
import json
import os
from pathlib import Path
import subprocess
import tempfile
import unittest
import zipfile
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas

SCRIPT = Path(__file__).resolve().parents[1] / 'Convert-DocumentsToJson.ps1'
PWSH = os.environ.get('PWSH', 'pwsh')


def package(path, entries):
    with zipfile.ZipFile(path, 'w') as archive:
        for name, xml in entries.items():
            archive.writestr(name, xml)


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.inputs = self.root / 'input'
        self.inputs.mkdir()
        self.output = self.root / 'output'

    def tearDown(self):
        self.temp.cleanup()

    def run_converter(self, *args, expected=0):
        backend_args = [] if '-PdfBackend' in args else ['-PdfBackend', os.environ.get('PDF_TEST_BACKEND', 'Poppler')]
        result = subprocess.run([PWSH, '-NoLogo', '-NoProfile', '-File', str(SCRIPT),
                                 '-InputPath', str(self.inputs), '-OutputDirectory', str(self.output),
                                 *backend_args, *args], capture_output=True, text=True)
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        return result

    def documents(self):
        result = {}
        for path in self.output.glob('*.json*'):
            raw = gzip.decompress(path.read_bytes()) if path.suffix == '.gz' else path.read_bytes()
            doc = json.loads(raw)
            result[doc['source']['name']] = doc
        return result

    def test_office_text_batch_and_repeat(self):
        (self.inputs / 'hello.txt').write_text('Café and Unicode: 世界', encoding='utf-8')
        package(self.inputs / 'word.docx', {'word/document.xml': '''
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:wp="urn:draw"><w:body>
<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Heading example</w:t></w:r></w:p>
<w:p><w:r><w:t>Hello </w:t></w:r><w:r><w:t>world</w:t></w:r><wp:docPr name="diagram" descr="Revenue chart"/></w:p>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Cell A</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Cell B</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
</w:body></w:document>'''})
        package(self.inputs / 'slides.pptx', {
            'ppt/presentation.xml': '<p:presentation xmlns:p="urn:p" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="1" r:id="r2"/><p:sldId id="2" r:id="r1"/></p:sldIdLst></p:presentation>',
            'ppt/_rels/presentation.xml.rels': '<Relationships><Relationship Id="r1" Target="slides/slide1.xml"/><Relationship Id="r2" Target="slides/slide2.xml"/></Relationships>',
            'ppt/slides/slide1.xml': '<sld><spTree><sp><p><r><t>Second slide</t></r></p></sp></spTree></sld>',
            'ppt/slides/slide2.xml': '<sld><spTree><sp><p><r><t>First slide</t></r></p></sp><pic><cNvPr descr="Chart"/></pic></spTree></sld>'})
        package(self.inputs / 'book.xlsx', {
            'xl/workbook.xml': '<workbook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Sales" r:id="r1"/></sheets></workbook>',
            'xl/_rels/workbook.xml.rels': '<Relationships><Relationship Id="r1" Target="worksheets/sheet1.xml"/></Relationships>',
            'xl/sharedStrings.xml': '<sst><si><t>Revenue</t></si></sst>',
            'xl/worksheets/sheet1.xml': '<worksheet><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c><c r="C1"><f>SUM(A2:A3)</f><v>42</v></c></row></sheetData></worksheet>'})
        self.run_converter()
        docs = self.documents()
        self.assertEqual(len(docs), 4)
        blocks = docs['word.docx']['units'][0]['blocks']
        self.assertEqual(blocks[0]['type'], 'heading')
        self.assertEqual(blocks[0]['level'], 1)
        self.assertEqual(blocks[1]['text'], 'Hello world')
        self.assertEqual(blocks[2]['alt_text'], 'Revenue chart')
        self.assertEqual(blocks[3]['rows'], [['Cell A', 'Cell B']])
        slides = docs['slides.pptx']['units']
        self.assertEqual([s['blocks'][0]['text'] for s in slides], ['First slide', 'Second slide'])
        cells = docs['book.xlsx']['units'][0]['blocks'][0]['rows'][0]['cells']
        self.assertEqual(cells[0]['value'], 'Revenue')
        self.assertEqual(cells[1]['cell'], 'C1')
        self.assertEqual(cells[1]['formula'], 'SUM(A2:A3)')
        self.assertEqual(docs['hello.txt']['units'][0]['blocks'][0]['text'], 'Café and Unicode: 世界')
        before = {p.name: p.read_bytes() for p in self.output.iterdir()}
        self.assertIn('SKIP', self.run_converter().stdout)
        self.assertEqual(before, {p.name: p.read_bytes() for p in self.output.iterdir()})
        (self.inputs / 'hello.txt').write_text('Updated')
        self.run_converter('-Force')
        self.assertEqual(self.documents()['hello.txt']['units'][0]['blocks'][0]['text'], 'Updated')

    def make_pdf(self):
        image = Image.new('RGB', (1800, 500), 'white')
        draw = ImageDraw.Draw(image)
        try:
            font = ImageFont.truetype('DejaVuSans.ttf', 70)
        except OSError:
            font = ImageFont.truetype('arial.ttf', 70)
        draw.text((60, 150), 'Scanned invoice total 12345', fill='black', font=font)
        scan = self.root / 'scan.png'
        image.save(scan)
        pdf = canvas.Canvas(str(self.inputs / 'mixed.pdf'))
        pdf.setFont('Helvetica', 20)
        pdf.drawString(60, 700, 'Native invoice total 67890')
        pdf.showPage()
        pdf.drawImage(str(scan), 25, 420, width=550, height=153)
        pdf.showPage()
        pdf.showPage()  # blank page must retain its page number
        pdf.save()

    def test_pdf_ocr_switch_auto_always_and_native(self):
        self.make_pdf()
        self.run_converter()
        units = self.documents()['mixed.pdf']['units']
        self.assertEqual(len(units), 3)
        self.assertEqual([u['number'] for u in units], [1, 2, 3])
        self.assertEqual(units[0]['extraction'], 'native')
        self.assertEqual(units[1]['blocks'], [])
        self.run_converter('-Ocr', '-Force')
        units = self.documents()['mixed.pdf']['units']
        self.assertEqual(units[0]['extraction'], 'native')
        self.assertIn('12345', units[1]['blocks'][0]['text'])
        self.run_converter('-Ocr', '-PdfOcr', 'Always', '-Force')
        units = self.documents()['mixed.pdf']['units']
        self.assertEqual(units[0]['extraction'], 'ocr')
        self.assertIn('67890', units[0]['blocks'][0]['text'])
        self.assertIn('12345', units[1]['blocks'][0]['text'])
        self.assertEqual(units[2]['blocks'], [])
        self.run_converter('-PdfOcr', 'Auto', '-Force')
        units = self.documents()['mixed.pdf']['units']
        self.assertEqual(units[0]['extraction'], 'native')
        self.assertEqual(units[1]['extraction'], 'ocr')
        self.run_converter('-PdfOcr', 'Never', '-Force')
        units = self.documents()['mixed.pdf']['units']
        self.assertEqual(units[1]['blocks'], [])

    def test_recursion_collisions_output_exclusion_and_failure_isolation(self):
        (self.inputs / 'a').mkdir()
        (self.inputs / 'b').mkdir()
        (self.inputs / 'a' / 'same.txt').write_text('First')
        (self.inputs / 'b' / 'same.txt').write_text('Second')
        (self.inputs / 'broken.pdf').write_text('Not a PDF')
        (self.inputs / 'legacy.doc').write_text('Old Word format')
        self.output = self.inputs / 'results'
        result = self.run_converter('-Recurse', '-PlainJson', expected=1)
        self.assertIn('Legacy .doc format', result.stdout)
        self.assertEqual(len(self.documents()), 1)  # names collide in dictionary only
        self.assertEqual(len(list(self.output.glob('*.json'))), 2)
        self.assertFalse(list(self.output.glob('*.tmp')))
        self.run_converter('-Recurse', '-PlainJson', expected=1)
        self.assertEqual(len(list(self.output.glob('*.json'))), 2)

    def test_real_office_packages(self):
        document = Document()
        document.add_heading('Quarterly results', level=1)
        document.add_paragraph('Revenue increased.')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Revenue'
        table.cell(0, 1).text = '12345'
        document.save(self.inputs / 'real.docx')
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = 'Quarterly presentation'
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        table.cell(0, 0).text = 'Revenue'
        table.cell(0, 1).text = '12345'
        presentation.save(self.inputs / 'real.pptx')
        workbook = Workbook()
        workbook.active.title = 'Quarterly data'
        workbook.active['A1'] = 'Revenue'
        workbook.active['C1'] = 12345
        workbook.save(self.inputs / 'real.xlsx')
        self.run_converter()
        docs = self.documents()
        self.assertEqual(docs['real.docx']['units'][0]['blocks'][0]['type'], 'heading')
        slide_blocks = docs['real.pptx']['units'][0]['blocks']
        self.assertEqual(slide_blocks[0]['type'], 'heading')
        self.assertEqual(slide_blocks[1]['rows'], [['Revenue', '12345']])
        sheet = docs['real.xlsx']['units'][0]
        self.assertEqual(sheet['name'], 'Quarterly data')
        cells = sheet['blocks'][0]['rows'][0]['cells']
        self.assertEqual(cells[0]['value'], 'Revenue')
        self.assertEqual(cells[1]['cell'], 'C1')

    @unittest.skipIf(os.name == 'nt', 'Linux guard only; requires Windows Acrobat for actual COM extraction')
    def test_acrobat_reports_windows_requirement(self):
        (self.inputs / 'sample.pdf').write_bytes(b'%PDF-1.4')
        result = self.run_converter('-PdfBackend', 'Acrobat', '-PdfOcr', 'Never', expected=1)
        self.assertIn('requires Windows', result.stdout)
        self.assertEqual(len(self.documents()), 0)


if __name__ == '__main__':
    unittest.main(verbosity=2)
